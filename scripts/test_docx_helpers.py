import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
table = doc.add_table(rows=2, cols=2)
cell = table.cell(0, 0)
shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
cell._tc.get_or_add_tcPr().append(shd)
print("Docx XML helpers work perfectly!")
