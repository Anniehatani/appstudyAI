import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Color Palette Constants (Nhóm TONE PINK & Thỏ Bảy Màu Kawaii Theme)
COLOR_BG = RGBColor(255, 247, 250)         # Soft Cream Blush #FFF7FA
COLOR_WHITE = RGBColor(255, 255, 255)      # Pure White #FFFFFF
COLOR_PRIMARY_PINK = RGBColor(255, 92, 138)# Vibrant Kawaii Pink #FF5C8A
COLOR_DEEP_BERRY = RGBColor(225, 29, 72)   # Deep Berry Rose #E11D48
COLOR_ROSE_DARK = RGBColor(159, 18, 57)    # Rose 900 #9F1239
COLOR_DARK_TEXT = RGBColor(30, 41, 59)     # Charcoal Slate #1E293B
COLOR_MUTED_TEXT = RGBColor(100, 116, 139) # Slate Muted #64748B
COLOR_CARD_BORDER = RGBColor(254, 205, 211)# Soft Pink Border #FECDD3
COLOR_CALLOUT_BG = RGBColor(255, 241, 242) # Callout Blush #FFF1F2
COLOR_ACCENT_YELLOW = RGBColor(245, 158, 11)# Gold Yellow #F59E0B
COLOR_ACCENT_PURPLE = RGBColor(139, 92, 246)# Soft Purple #8B5CF6
COLOR_ACCENT_GREEN = RGBColor(16, 185, 129)# Mint Green #10B981
COLOR_NAVY = RGBColor(30, 58, 138)         # Deep Navy #1E3A8A

def add_full_background(slide):
    """Draws full slide soft blush background with rainbow accent strip at top."""
    # Base background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()

    # Top Rainbow Accent Ribbon (Thỏ Bảy Màu signature)
    colors = [
        RGBColor(239, 68, 68),   # Red
        RGBColor(249, 115, 22),  # Orange
        RGBColor(234, 179, 8),   # Yellow
        RGBColor(34, 197, 94),   # Green
        RGBColor(6, 182, 212),   # Cyan
        RGBColor(59, 130, 246),  # Blue
        RGBColor(168, 85, 247)   # Purple
    ]
    seg_w = 13.333 / 7.0
    for idx, c in enumerate(colors):
        ribbon = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(idx * seg_w), 0, Inches(seg_w + 0.05), Inches(0.08))
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = c
        ribbon.line.fill.background()

def add_slide_header(slide, title, category="TỔNG KẾT SPRINT 2 • NHÓM TONE PINK • GV: VÂN ANH", slide_num=1, total_slides=15):
    """Creates a modern header with category chip, title and slide counter."""
    # Category Chip
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.28), Inches(4.8), Inches(0.32))
    chip.fill.solid()
    chip.fill.fore_color.rgb = RGBColor(255, 228, 235)
    chip.line.color.rgb = COLOR_PRIMARY_PINK
    chip.line.width = Pt(1)
    tf_c = chip.text_frame
    tf_c.word_wrap = False
    p_c = tf_c.paragraphs[0]
    p_c.text = f"🌸 {category}"
    p_c.font.name = "Arial"
    p_c.font.size = Pt(9.5)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_DEEP_BERRY

    # Main Title
    tb_tit = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(9.8), Inches(0.65))
    tf_t = tb_tit.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title
    p_t.font.name = "Arial"
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_DARK_TEXT

    # Counter Pill at Top Right
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.4), Inches(0.35), Inches(1.15), Inches(0.35))
    pill.fill.solid()
    pill.fill.fore_color.rgb = COLOR_WHITE
    pill.line.color.rgb = COLOR_CARD_BORDER
    pill.line.width = Pt(1)
    tf_p = pill.text_frame
    p_p = tf_p.paragraphs[0]
    p_p.alignment = PP_ALIGN.CENTER
    p_p.text = f"{slide_num:02d} / {total_slides:02d}"
    p_p.font.name = "Arial"
    p_p.font.size = Pt(10)
    p_p.font.bold = True
    p_p.font.color.rgb = COLOR_DEEP_BERRY

def add_card(slide, left, top, width, height, title="", border_color=COLOR_CARD_BORDER, bg_color=COLOR_WHITE):
    """Draws a rounded white card with subtle border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    if title:
        # Title banner on top of card
        tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.12), Inches(width - 0.3), Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_DEEP_BERRY
    return card

def create_sprint2_slides():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: TRANG BÌA ====================
    s1 = prs.slides.add_slide(blank_layout)
    add_full_background(s1)

    # Big Card Center
    add_card(s1, 1.2, 0.8, 10.933, 5.9, border_color=COLOR_PRIMARY_PINK, bg_color=COLOR_WHITE)

    # University & Department
    tb1_u = s1.shapes.add_textbox(Inches(1.5), Inches(1.0), Inches(10.3), Inches(0.6))
    p1_u = tb1_u.text_frame.paragraphs[0]
    p1_u.alignment = PP_ALIGN.CENTER
    p1_u.text = "TRƯỜNG ĐẠI HỌC CMC • KHOA CÔNG NGHỆ THÔNG TIN"
    p1_u.font.name = "Arial"
    p1_u.font.size = Pt(12)
    p1_u.font.bold = True
    p1_u.font.color.rgb = COLOR_MUTED_TEXT

    # Course Badge
    badge1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.5), Inches(4.333), Inches(0.42))
    badge1.fill.solid()
    badge1.fill.fore_color.rgb = RGBColor(255, 228, 235)
    badge1.line.color.rgb = COLOR_PRIMARY_PINK
    p1_b = badge1.text_frame.paragraphs[0]
    p1_b.alignment = PP_ALIGN.CENTER
    p1_b.text = "MÔN: ỨNG DỤNG PHẦN MỀM"
    p1_b.font.name = "Arial"
    p1_b.font.size = Pt(11)
    p1_b.font.bold = True
    p1_b.font.color.rgb = COLOR_DEEP_BERRY

    img_icon = os.path.abspath("tổng kết sprint 2/assets/app_icon.png")
    if os.path.exists(img_icon):
        s1.shapes.add_picture(img_icon, Inches(1.6), Inches(1.15), width=Inches(0.9), height=Inches(0.9))
        s1.shapes.add_picture(img_icon, Inches(10.8), Inches(1.15), width=Inches(0.9), height=Inches(0.9))

    # Main Title
    tb1_t = s1.shapes.add_textbox(Inches(1.5), Inches(2.05), Inches(10.3), Inches(1.2))
    p1_t = tb1_t.text_frame.paragraphs[0]
    p1_t.alignment = PP_ALIGN.CENTER
    p1_t.text = "BÁO CÁO TỔNG KẾT SPRINT 2"
    p1_t.font.name = "Arial"
    p1_t.font.size = Pt(32)
    p1_t.font.bold = True
    p1_t.font.color.rgb = COLOR_DEEP_BERRY

    # Subtitle
    tb1_s = s1.shapes.add_textbox(Inches(1.5), Inches(2.95), Inches(10.3), Inches(0.8))
    p1_s = tb1_s.text_frame.paragraphs[0]
    p1_s.alignment = PP_ALIGN.CENTER
    p1_s.text = "DỰ ÁN 3TG ALL-IN-ONE: MẠNG XÃ HỘI HỌC TẬP, GAME HÓA CHUỖI STREAK DUOLINGO,\nBỘ SƯU TẬP 20 MASCOT THỎ BẢY MÀU & ĐÓNG GÓI DESKTOP ELECTRON"
    p1_s.font.name = "Arial"
    p1_s.font.size = Pt(13)
    p1_s.font.bold = True
    p1_s.font.color.rgb = COLOR_NAVY

    # Info Box: Lecturer & Team
    info_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(3.9), Inches(8.933), Inches(2.4))
    info_box.fill.solid()
    info_box.fill.fore_color.rgb = COLOR_CALLOUT_BG
    info_box.line.color.rgb = COLOR_CARD_BORDER
    tf_ib = info_box.text_frame
    tf_ib.word_wrap = True

    p_gv = tf_ib.paragraphs[0]
    p_gv.alignment = PP_ALIGN.CENTER
    p_gv.text = "🌸 Giảng viên hướng dẫn: CÔ VÂN ANH"
    p_gv.font.name = "Arial"
    p_gv.font.size = Pt(13)
    p_gv.font.bold = True
    p_gv.font.color.rgb = COLOR_DEEP_BERRY

    p_group = tf_ib.add_paragraph()
    p_group.alignment = PP_ALIGN.CENTER
    p_group.text = "Nhóm thực hiện: NHÓM TONE PINK"
    p_group.font.name = "Arial"
    p_group.font.size = Pt(13.5)
    p_group.font.bold = True
    p_group.font.color.rgb = COLOR_NAVY

    p_m1 = tf_ib.add_paragraph()
    p_m1.alignment = PP_ALIGN.CENTER
    p_m1.text = "1. Lê Văn An (Trưởng nhóm — Kiến trúc Cloud D1/R2, Mạng xã hội & Đóng gói Desktop)"
    p_m1.font.name = "Arial"
    p_m1.font.size = Pt(11)
    p_m1.font.color.rgb = COLOR_DARK_TEXT

    p_m2 = tf_ib.add_paragraph()
    p_m2.alignment = PP_ALIGN.CENTER
    p_m2.text = "2. Nguyễn Như Quỳnh (UI/UX Tone Pink, Mascot 20 Biểu Cảm & Ôn Tập Câu Sai)"
    p_m2.font.name = "Arial"
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = COLOR_DARK_TEXT

    p_m3 = tf_ib.add_paragraph()
    p_m3.alignment = PP_ALIGN.CENTER
    p_m3.text = "3. Nguyễn Văn Tú (Logic Chuỗi Streak Duolingo, Nhiệm Vụ Ngày, Giải Đấu Tuần & Test)"
    p_m3.font.name = "Arial"
    p_m3.font.size = Pt(11)
    p_m3.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 2: MỤC TIÊU & NHÂN SỰ ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_full_background(s2)
    add_slide_header(s2, "Mục Tiêu Trọng Tâm Sprint 2 & Phân Công Nhân Sự", slide_num=2)

    # Left Card: 4 Trụ Cột
    add_card(s2, 0.8, 1.4, 5.7, 5.5, "🎯 4 Trụ Cột Chiến Lược Sprint 2")
    tb2_l = s2.shapes.add_textbox(Inches(0.95), Inches(1.95), Inches(5.4), Inches(4.8))
    tf2_l = tb2_l.text_frame
    tf2_l.word_wrap = True

    pts_pillar = [
        ("Mạng xã hội học tập (Social Learning):", "Soi hồ sơ bạn học khác, Kết bạn 1-chạm, Nhắn tin trực tiếp thời gian thực có unread badge."),
        ("Game hóa sâu sắc (Gamification):", "Chuỗi ngày học Duolingo với cơ chế Hóa đá 2 lần/tuần, Nhiệm vụ ngày và Giải đấu tuần thưởng Kim Cương."),
        ("Tái cấu trúc Ôn tập câu sai:", "Phân môn rành mạch, làm lại trắc nghiệm 4 đáp án trực tiếp, sửa triệt để lỗi kiểm tra đáp án bị treo."),
        ("Độc lập & Trải nghiệm Desktop:", "Custom Dropdown 3D Tone Pink, Tối ưu màn hình Desktop tỷ lệ 150%, Đóng gói file exe cài đặt độc lập.")
    ]
    for idx, (head, body) in enumerate(pts_pillar):
        p = tf2_l.paragraphs[0] if idx == 0 else tf2_l.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {head} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_DEEP_BERRY
        r2 = p.add_run()
        r2.text = f"{body}\n"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_DARK_TEXT

    # Right Card: 3 Thành Viên
    add_card(s2, 6.8, 1.4, 5.7, 5.5, "👥 Đội Ngũ Thực Hiện (Nhóm TONE PINK)")
    tb2_r = s2.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.4), Inches(4.8))
    tf2_r = tb2_r.text_frame
    tf2_r.word_wrap = True

    pts_members = [
        ("Lê Văn An — Trưởng nhóm (Lead Architect)", "• Quản lý tiến độ toàn bộ Sprint 2\n• Thiết kế API D1/R2 Mạng xã hội, Kết bạn & Nhắn tin\n• Tối ưu hệ thống & Đóng gói Electron Exe Standalone\n• Đánh giá nỗ lực: 100% Hoàn thành xuất sắc"),
        ("Nguyễn Như Quỳnh — Frontend & UI/UX Designer", "• Thiết kế giao diện Tone Pink chuẩn Thỏ Bảy Màu\n• Tái cấu trúc phân hệ Ôn tập câu sai trắc nghiệm theo môn\n• Sáng tạo phòng thay đồ 20 Mascot Thỏ Bảy Màu\n• Đánh giá nỗ lực: 100% Hoàn thành xuất sắc"),
        ("Nguyễn Văn Tú — Backend & Quality Assurance", "• Lập trình thuật toán Chuỗi Streak Freeze Duolingo\n• Xây dựng hệ thống Nhiệm vụ ngày & Giải đấu tuần\n• Thực thi bộ kịch bản kiểm thử tự động, build NSIS\n• Đánh giá nỗ lực: 100% Hoàn thành xuất sắc")
    ]
    for idx, (head, body) in enumerate(pts_members):
        p = tf2_r.paragraphs[0] if idx == 0 else tf2_r.add_paragraph()
        r1 = p.add_run()
        r1.text = f"🌟 {head}\n"
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = COLOR_NAVY
        r2 = p.add_run()
        r2.text = f"{body}\n\n"
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 3: TỔNG QUAN KIẾN TRÚC ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_full_background(s3)
    add_slide_header(s3, "Tổng Quan Hệ Thống Tính Năng Đột Phá Sprint 2", slide_num=3)

    cards_s3 = [
        (0.8, 1.4, 3.6, 2.6, "1. Mạng Xã Hội Học Tập", "• Soi Profile chi tiết bạn học\n• Kết bạn 1-chạm thông minh\n• Chat Real-Time & Unread Badge\n• Empty state vui nhộn với Mascot", COLOR_PRIMARY_PINK),
        (4.8, 1.4, 3.6, 2.6, "2. Game Hóa Chuỗi Duolingo", "• Giữ chuỗi đóng băng (2 lần/tuần)\n• Biểu tượng hóa đá 🧊 trên lịch\n• Tự động chuyển đổi lịch thường\n• Cửa hàng mua Bảo vệ chuỗi", COLOR_DEEP_BERRY),
        (8.8, 1.4, 3.6, 2.6, "3. Bộ Sưu Tập 20 Mascot", "• 6 Mascot cơ bản mặc định\n• Gói Học Bá Siêu Đẳng (7 con xịn)\n• Gói Thánh Hề Bá Đạo (7 con meme)\n• Mascot Wardrobe phân loại gọn", COLOR_ACCENT_PURPLE),
        (0.8, 4.3, 3.6, 2.6, "4. Ôn Tập Câu Sai Theo Môn", "• Phân môn riêng biệt chống lộn\n• Làm lại trắc nghiệm 4 đáp án\n• Phản hồi tức thì & Lời giải chi tiết\n• Hồi phục Tim ❤️ & Thưởng XP", COLOR_ACCENT_YELLOW),
        (4.8, 4.3, 3.6, 2.6, "5. Quests & Weekly Leagues", "• Nhiệm vụ ngày đa dạng thưởng 💎\n• Giải đấu 5 hạng (Đồng - Kim Cương)\n• Top 1 nhận 100 💎, Top 2 60 💎\n• Modal nhận thưởng vinh quang", COLOR_ACCENT_GREEN),
        (8.8, 4.3, 3.6, 2.6, "6. Desktop Exe Standalone", "• Đóng gói Setup.exe & Portable\n• Nhúng máy chủ embedded-server\n• Local Replica SQLite ngoại tuyến\n• Tối ưu màn hình 150% Scaling", COLOR_NAVY)
    ]
    for (cl, ct, cw, ch, c_tit, c_body, c_color) in cards_s3:
        add_card(s3, cl, ct, cw, ch, c_tit, border_color=c_color)
        tb = s3.shapes.add_textbox(Inches(cl + 0.15), Inches(ct + 0.55), Inches(cw - 0.3), Inches(ch - 0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = c_body
        p.font.name = "Arial"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 4: SOI HỒ SƠ & BẠN BÈ ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_full_background(s4)
    add_slide_header(s4, "Mạng Xã Hội Học Tập — Soi Hồ Sơ & Kết Bạn 1-Chạm", slide_num=4)

    add_card(s4, 0.8, 1.4, 6.0, 5.5, "🔍 Cơ Chế Soi Hồ Sơ (Profile Inspection)")
    tb4_l = s4.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.7), Inches(4.7))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True
    p4_1 = tf4_l.paragraphs[0]
    p4_1.text = "• Endpoint chuyên dụng: /api/users/:id\n\n" \
                "• Thông tin hiển thị đầy đủ, chân thực:\n" \
                "   - Ảnh đại diện Avatar & Ảnh bìa Cover từ Cloudflare R2\n" \
                "   - Họ tên, Trường đại học, Chuyên ngành, Tiểu sử Bio\n" \
                "   - Tổng số bài học đã hoàn thành\n" \
                "   - Tổng điểm kinh nghiệm XP tích lũy\n" \
                "   - Chuỗi ngày học liên tục (Streak) hiện tại\n" \
                "   - Thứ hạng và danh hiệu Giải đấu tuần\n\n" \
                "• Truy cập tức thì: Học viên chỉ cần bấm vào tên hoặc avatar của bất kỳ ai trên Bảng xếp hạng hoặc danh sách bạn bè để 'soi' profile chi tiết."
    p4_1.font.name = "Arial"
    p4_1.font.size = Pt(11)
    p4_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s4, 7.1, 1.4, 5.4, 5.5, "🤝 Vòng Đời Kết Bạn (Friendship Lifecycle)")
    tb4_r = s4.shapes.add_textbox(Inches(7.25), Inches(2.0), Inches(5.1), Inches(4.7))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True
    p4_2 = tf4_r.paragraphs[0]
    p4_2.text = "• Bảng dữ liệu D1: friendships\n" \
                "   - Cấu trúc: id, user_id, friend_id, status, created_at\n" \
                "   - Ràng buộc: UNIQUE(user_id, friend_id) chống gửi trùng\n\n" \
                "• Các trạng thái kết bạn:\n" \
                "   1. Chưa kết bạn: Nút '+ Kết bạn' màu hồng nổi bật\n" \
                "   2. Đã gửi lời mời: Nút 'Đang chờ' tinh tế\n" \
                "   3. Đã là bạn bè: Nút 'Bạn bè ✓' kèm nút 'Nhắn tin'\n\n" \
                "• Danh sách bạn bè sống động: Hiển thị danh sách kèm trạng thái online/offline, xem nhanh cấp bậc giải đấu của từng người bạn."
    p4_2.font.name = "Arial"
    p4_2.font.size = Pt(11)
    p4_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 5: NHẮN TIN TRỰC TIẾP ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_full_background(s5)
    add_slide_header(s5, "Nhắn Tin Trực Tiếp Thời Gian Thực (Direct Messaging)", slide_num=5)

    add_card(s5, 0.8, 1.4, 5.7, 5.5, "💬 Giao Diện & Trải Nghiệm Trò Chuyện")
    tb5_l = s5.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf5_l = tb5_l.text_frame
    tf5_l.word_wrap = True
    p5_1 = tf5_l.paragraphs[0]
    p5_1.text = "• Đưa tin nhắn làm trọng tâm chính: Khi truy cập mục Bạn bè, khung trò chuyện mở ngay lập tức mà không cần qua nhiều bước trung gian.\n\n" \
                "• Thiết kế Bong bóng Chat hiện đại:\n" \
                "   - Tin nhắn của mình (Mine): Nền Hồng Gradient tươi sáng, căn lề phải, bo góc chuẩn mực.\n" \
                "   - Tin nhắn của bạn (Other): Nền xám thanh lịch, căn lề trái.\n" \
                "   - Dấu mốc thời gian: Hiển thị giờ gửi rõ ràng dưới mỗi bong bóng chat.\n\n" \
                "• Tự động cuộn: Luôn tự động cuộn xuống tin nhắn mới nhất khi có tin đến."
    p5_1.font.name = "Arial"
    p5_1.font.size = Pt(11)
    p5_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s5, 6.8, 1.4, 5.7, 5.5, "📸 Giao Diện Nhắn Tin Thực Tế")
    tb5_r = s5.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.4), Inches(1.4))
    tf5_r = tb5_r.text_frame
    tf5_r.word_wrap = True
    p5_2 = tf5_r.paragraphs[0]
    p5_2.text = "• Chấm đỏ unread badge nhấp nháy ở Sidebar và Tab\n" \
                "• Bong bóng chat Mine (Hồng Gradient) vs Other (Xám)\n" \
                "• Empty state vui nhộn cùng Mascot Thỏ Bảy Màu"
    p5_2.font.name = "Arial"
    p5_2.font.size = Pt(10)
    p5_2.font.color.rgb = COLOR_DEEP_BERRY

    img_dc = os.path.abspath("tổng kết sprint 2/assets/direct_chat.png")
    if os.path.exists(img_dc):
        s5.shapes.add_picture(img_dc, Inches(8.3), Inches(3.3), width=Inches(2.7))

    # ==================== SLIDE 6: ÔN TẬP CÂU SAI ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_full_background(s6)
    add_slide_header(s6, "Tái Cấu Trúc Toàn Diện Phân Hệ Ôn Tập Câu Sai", slide_num=6)

    add_card(s6, 0.8, 1.4, 5.7, 5.5, "📚 Phân Môn & Trắc Nghiệm Tức Thì")
    tb6_l = s6.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf6_l = tb6_l.text_frame
    tf6_l.word_wrap = True
    p6_1 = tf6_l.paragraphs[0]
    p6_1.text = "• Phân loại câu sai theo từng môn học riêng biệt:\n" \
                "   - Sinh viên tự chọn môn muốn ôn tập (Kỹ năng mềm, Lập trình C++, Toán rời rạc, Web, v.v.).\n" \
                "   - Ngăn ngừa triệt để việc lẫn lộn câu sai giữa các môn học khác nhau.\n\n" \
                "• Làm lại trực tiếp trắc nghiệm 4 đáp án (A, B, C, D):\n" \
                "   - Thay vì chỉ hỏi nhớ hay chưa, học viên phải tư duy giải lại câu hỏi và chọn đáp án.\n" \
                "   - Giao diện thẻ đáp án lớn, hiệu ứng hover nhấp nhô sống động chuẩn Duolingo."
    p6_1.font.name = "Arial"
    p6_1.font.size = Pt(11)
    p6_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s6, 6.8, 1.4, 5.7, 5.5, "📸 Giao Diện Làm Lại Câu Sai Trắc Nghiệm")
    tb6_r = s6.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.4), Inches(1.3))
    tf6_r = tb6_r.text_frame
    tf6_r.word_wrap = True
    p6_2 = tf6_r.paragraphs[0]
    p6_2.text = "• Sửa triệt để lỗi kiểm tra bị treo, phản hồi < 50ms\n" \
                "• Lời giải chi tiết (Explanation) giải thích bản chất\n" \
                "• Khắc phục đúng: Hồi phục 1 Tim ❤️ và cộng 5 XP"
    p6_2.font.name = "Arial"
    p6_2.font.size = Pt(10)
    p6_2.font.color.rgb = COLOR_DEEP_BERRY

    img_mr = os.path.abspath("tổng kết sprint 2/assets/mistakes_review.png")
    if os.path.exists(img_mr):
        s6.shapes.add_picture(img_mr, Inches(7.1), Inches(3.3), width=Inches(5.1))

    # ==================== SLIDE 7: CHUỖI DUOLINGO & STREAK FREEZE ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_full_background(s7)
    add_slide_header(s7, "Chuỗi Học Tập Duolingo & Cơ Chế Hóa Đá (Streak Freeze)", slide_num=7)

    add_card(s7, 0.8, 1.4, 5.7, 5.5, "🧊 Cơ Chế Hóa Đá Chuỗi (Streak Freeze)")
    tb7_l = s7.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True
    p7_1 = tf7_l.paragraphs[0]
    p7_1.text = "• Chuẩn mực Duolingo Quốc tế:\n" \
                "   - Mỗi tuần học viên được cấp 2 lượt Bảo vệ chuỗi (Streak Freeze).\n" \
                "   - Khi có ngày bận không học, chuỗi ngày không bị đứt mà được 'đóng băng an toàn'.\n\n" \
                "• Hiệu ứng Tảng Băng Hóa Đá:\n" \
                "   - Trên lịch ngày học, ngày được bảo vệ sẽ hiển thị biểu tượng tảng băng 🧊 với viền xanh băng tuyết lấp lánh.\n\n" \
                "• Tự động rã đông khi học bài mới:\n" \
                "   - Chỉ cần hoàn thành 1 bài học bất kỳ, chuỗi được rã đông và tiếp tục tích lũy tăng đều."
    p7_1.font.name = "Arial"
    p7_1.font.size = Pt(11)
    p7_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s7, 6.8, 1.4, 5.7, 5.5, "📸 Giao Diện Lịch Duolingo Hóa Đá 🧊")
    tb7_r = s7.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.4), Inches(1.4))
    tf7_r = tb7_r.text_frame
    tf7_r.word_wrap = True
    p7_2 = tf7_r.paragraphs[0]
    p7_2.text = "• Biểu tượng tảng băng 🧊 xuất hiện khi dùng freeze\n" \
                "• Quá 2 ngày: Tự động chuyển về lịch số bình thường\n" \
                "• Mua vật phẩm 'Bảo vệ chuỗi' tại Cửa hàng bằng 💎"
    p7_2.font.name = "Arial"
    p7_2.font.size = Pt(10)
    p7_2.font.color.rgb = COLOR_DEEP_BERRY

    img_sf = os.path.abspath("tổng kết sprint 2/assets/streak_freeze.png")
    if os.path.exists(img_sf):
        s7.shapes.add_picture(img_sf, Inches(7.1), Inches(3.4), width=Inches(5.1))
    p7_2.font.name = "Arial"
    p7_2.font.size = Pt(11)
    p7_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 8: 20 MASCOT THỎ BẢY MÀU ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_full_background(s8)
    add_slide_header(s8, "Hệ Sinh Thái 20 Mascot Thỏ Bảy Màu & Gói Mở Khóa Gacha", slide_num=8)

    cards_s8 = [
        (0.8, 1.4, 3.6, 5.5, "🐰 6 Mascot Cơ Bản\n(Mặc định miễn phí)",
         "• Dành cho mọi tân thủ bắt đầu hành trình học tập:\n\n"
         "1. Thỏ Chăm Chỉ (Cầm bút ghi chép)\n"
         "2. Thỏ Suy Tư (Gãi đầu suy nghĩ bài)\n"
         "3. Thỏ Vui Vẻ (Cười tươi vẫy tay chào)\n"
         "4. Thỏ Ngủ Gật (Mắt lim dim ôm gối)\n"
         "5. Thỏ Quyết Tâm (Thắt băng đô đỏ)\n"
         "6. Thỏ Đọc Sách (Đeo kính tri thức)\n\n"
         "👉 Sẵn sàng sử dụng ngay sau khi đăng ký tài khoản 3TG!"),

        (4.8, 1.4, 3.6, 5.5, "🎓 Gói Học Bá Siêu Đẳng\n(7 Mascot Hào Quang)",
         "• Mở khóa bằng Kim Cương học tập:\n\n"
         "7. Thỏ Tiến Sĩ (Đội mũ cử nhân)\n"
         "8. Thỏ Bác Học (Cầm ống nghiệm hóa chất)\n"
         "9. Thỏ Lập Trình Viên (Ôm laptop gõ code)\n"
         "10. Thỏ Thủ Khoa (Đeo vòng nguyệt quế)\n"
         "11. Thỏ Thần Đồng (Hào quang vàng rực)\n"
         "12. Thỏ Đạt Giải Nhất (Cầm cúp vàng 🏆)\n"
         "13. Thỏ Giáo Sư (Cầm phấn giảng bài)\n\n"
         "👉 Sở hữu hiệu ứng hào quang lấp lánh!"),

        (8.8, 1.4, 3.6, 5.5, "🎭 Gói Thánh Hề Bá Đạo\n(7 Mascot Hài Hước)",
         "• Bộ sưu tập Meme Thỏ Bảy Màu độc lạ:\n\n"
         "14. Thỏ Troll Face (Nụ cười thương hiệu)\n"
         "15. Thỏ Cà Khịa (Chống nạnh cà khịa bạn)\n"
         "16. Thỏ Điệp Viên (Đeo kính đen ngầu lòi)\n"
         "17. Thỏ Siêu Nhân (Áo choàng bay phấp phới)\n"
         "18. Thỏ Ninja (Bịt mặt bí ẩn)\n"
         "19. Thỏ Bốc Hỏa (Cháy hết mình ôn thi)\n"
         "20. Thỏ Đại Gia (Đeo dây chuyền vàng)\n\n"
         "👉 Hoạt ảnh động độc quyền vui nhộn!")
    ]
    for (cl, ct, cw, ch, c_tit, c_body) in cards_s8:
        add_card(s8, cl, ct, cw, ch, c_tit, border_color=COLOR_PRIMARY_PINK)
        tb = s8.shapes.add_textbox(Inches(cl + 0.15), Inches(ct + 0.9), Inches(cw - 0.3), Inches(ch - 1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = c_body
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 9: QUESTS & LEAGUES ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_full_background(s9)
    add_slide_header(s9, "Nhiệm Vụ Hàng Ngày (Quests) & Giải Đấu Tuần Thưởng Kim Cương", slide_num=9)

    add_card(s9, 0.8, 1.4, 5.7, 5.5, "📜 Nhiệm Vụ Hàng Ngày (Daily Quests)")
    tb9_l = s9.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    p9_1 = tf9_l.paragraphs[0]
    p9_1.text = "• Hệ thống tự động làm mới vào 00:00 mỗi ngày:\n\n" \
                "1. 🌅 Khởi động ngày mới: Học bài đầu tiên (Thưởng: 10 💎)\n" \
                "2. 📚 Chăm chỉ rèn luyện: Hoàn thành 3 bài học (Thưởng: 15 💎)\n" \
                "3. ⚡ Chiến thần học tập: Hoàn thành 5 bài học (Thưởng: 20 💎)\n" \
                "4. 🎯 Điểm số hoàn hảo: Đạt điểm bài học ≥ 80% (Thưởng: 15 💎)\n" \
                "5. 🩹 Khắc phục điểm yếu: Ôn tập 1 câu sai (Thưởng: 10 💎)\n" \
                "6. 💎 Tích lũy tri thức: Đạt 30 XP trong ngày (Thưởng: 15 💎)\n\n" \
                "• Nút 'Nhận Thưởng' xanh lá nổi bật với hiệu ứng âm thanh fanfare vui nhộn!"
    p9_1.font.name = "Arial"
    p9_1.font.size = Pt(10.5)
    p9_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s9, 6.8, 1.4, 5.7, 5.5, "📸 Giao Diện Nhiệm Vụ Ngày & Giải Đấu 🏆")
    tb9_r = s9.shapes.add_textbox(Inches(6.95), Inches(1.95), Inches(5.4), Inches(1.3))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    p9_2 = tf9_r.paragraphs[0]
    p9_2.text = "• 5 Bậc giải đấu danh giá: Đồng → Bạc → Vàng → Lục Bảo → Kim Cương\n" \
                "• Top 1 nhận 100 💎, Top 2 60 💎, Top 3 40 💎 mỗi tuần\n" \
                "• Nút nhận thưởng Kim Cương nổi bật có âm thanh vui nhộn"
    p9_2.font.name = "Arial"
    p9_2.font.size = Pt(9.5)
    p9_2.font.color.rgb = COLOR_DEEP_BERRY

    img_dq = os.path.abspath("tổng kết sprint 2/assets/daily_quests.png")
    if os.path.exists(img_dq):
        s9.shapes.add_picture(img_dq, Inches(7.3), Inches(3.3), width=Inches(4.7))

    # ==================== SLIDE 10: CUSTOM DROPDOWN & AI CHAT ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_full_background(s10)
    add_slide_header(s10, "Custom Dropdown 3D Tone Pink & Trợ Giảng AI Thông Minh", slide_num=10)

    add_card(s10, 0.8, 1.4, 5.7, 5.5, "🎨 Custom Dropdown Card 3D Tone Pink")
    tb10_l = s10.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf10_l = tb10_l.text_frame
    tf10_l.word_wrap = True
    p10_1 = tf10_l.paragraphs[0]
    p10_1.text = "• Loại bỏ thẻ select Windows cũ kỹ:\n" \
                 "   - Thẻ select mặc định trên Windows hiển thị đơn điệu, không phù hợp thiết kế Tone Pink.\n\n" \
                 "• Sáng tạo linh kiện lumi-dropdown-panel:\n" \
                 "   - Nút kích hoạt dạng Card 3D màu hồng pastel mềm mại.\n" \
                 "   - Menu xổ xuống với hiệu ứng kính mờ (blur backdrop glassmorphism) sang trọng.\n" \
                 "   - Danh sách lịch sử chat AI được gom nhóm theo môn học hoặc hỏi đáp tự do.\n" \
                 "   - Nút '+ Tạo cuộc trò chuyện mới' luôn nổi bật ở đỉnh menu."
    p10_1.font.name = "Arial"
    p10_1.font.size = Pt(11)
    p10_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s10, 6.8, 1.4, 5.7, 5.5, "🤖 Trợ Giảng AI Lumi & Giáo Sư Thỏ")
    tb10_r = s10.shapes.add_textbox(Inches(6.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf10_r = tb10_r.text_frame
    tf10_r.word_wrap = True
    p10_2 = tf10_r.paragraphs[0]
    p10_2.text = "• Kết nối trực tiếp mô hình DeepSeek AI:\n" \
                 "   - Khả năng suy luận lập trình, toán học và kỹ năng mềm xuất sắc.\n\n" \
                 "• Công nghệ trích xuất RAG (Retrieval-Augmented Generation):\n" \
                 "   - Tự động trích dẫn slide bài giảng, giáo trình chính thống của Đại học CMC.\n" \
                 "   - Câu trả lời chuẩn xác theo đề cương môn học thực tế, không bịa đặt kiến thức.\n\n" \
                 "• Phong cách hướng dẫn:\n" \
                 "   - Hóa thân thành Giáo sư Thỏ Bảy Màu hóm hỉnh, ân cần, giải thích từ cơ bản đến nâng cao."
    p10_2.font.name = "Arial"
    p10_2.font.size = Pt(11)
    p10_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 11: RESPONSIVE 150% SCALING ====================
    s11 = prs.slides.add_slide(blank_layout)
    add_full_background(s11)
    add_slide_header(s11, "Tối Ưu Màn Hình Desktop 150% Scaling & Web Mobile", slide_num=11)

    add_card(s11, 0.8, 1.4, 5.7, 5.5, "🖥️ Tối Ưu Desktop Tỷ Lệ 150% (Windows Scaling)")
    tb11_l = s11.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf11_l = tb11_l.text_frame
    tf11_l.word_wrap = True
    p11_1 = tf11_l.paragraphs[0]
    p11_1.text = "• Thách thức thực tế:\n" \
                 "   - Đa số laptop sinh viên chạy Windows 10/11 có màn hình độ phân giải cao được hệ điều hành tự động đặt tỷ lệ hiển thị Display Scale 125% - 150%.\n" \
                 "   - Ở phiên bản trước, giao diện bị khuất lề phải và che một phần Sidebar bên trái.\n\n" \
                 "• Giải pháp từ nhóm TONE PINK:\n" \
                 "   - Thiết kế lại hệ thống layout dựa trên Viewport Dynamic Units (dvh/cqw).\n" \
                 "   - Sidebar được co giãn tỷ lệ vàng (240px $\rightarrow$ 260px) tự động phân bổ đều các nút điều hướng.\n" \
                 "   - Chống tràn 100% trên mọi tỷ lệ màn hình máy tính."
    p11_1.font.name = "Arial"
    p11_1.font.size = Pt(11)
    p11_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s11, 6.8, 1.4, 5.7, 5.5, "📱 Trải Nghiệm Mượt Mà Trên Web Mobile")
    tb11_r = s11.shapes.add_textbox(Inches(6.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf11_r = tb11_r.text_frame
    tf11_r.word_wrap = True
    p11_2 = tf11_r.paragraphs[0]
    p11_2.text = "• Thanh điều hướng Mobile Bottom Nav:\n" \
                 "   - Khi xem trên điện thoại, thanh Sidebar tự động chuyển thành Bottom Navigation Bar 5 nút tiện lợi nằm ở đáy màn hình.\n\n" \
                 "• Touch Friendly UI:\n" \
                 "   - Các nút chọn đáp án, nút kiểm tra và menu tin nhắn có kích thước tối thiểu 44px $\times$ 44px theo chuẩn thiết kế Apple iOS HIG & Android Material 3.\n\n" \
                 "• Tốc độ tải trang siêu tốc:\n" \
                 "   - Tối ưu hóa dung lượng CSS/JS, tải trang trong 0.8 giây trên mạng 4G/5G."
    p11_2.font.name = "Arial"
    p11_2.font.size = Pt(11)
    p11_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 12: ĐÓNG GÓI ELECTRON DESKTOP ====================
    s12 = prs.slides.add_slide(blank_layout)
    add_full_background(s12)
    add_slide_header(s12, "Đóng Gói Ứng Dụng Desktop Độc Lập (Electron Standalone)", slide_num=12)

    add_card(s12, 0.8, 1.4, 5.7, 5.5, "📦 Hai Định Dạng Phân Phối Tiện Lợi")
    tb12_l = s12.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf12_l = tb12_l.text_frame
    tf12_l.word_wrap = True
    p12_1 = tf12_l.paragraphs[0]
    p12_1.text = "1. Bản Cài Đặt NSIS Wizard:\n" \
                 "   - File: 3TG Học Vui Mỗi Ngày Setup 1.0.0.exe\n" \
                 "   - Trình cài đặt đồ họa chuyên nghiệp, cho phép tùy chọn thư mục cài đặt, tự động tạo Icon ngoài Desktop và Start Menu.\n\n" \
                 "2. Bản Di Động Portable:\n" \
                 "   - File: 3TG Học Vui Mỗi Ngày 1.0.0.exe\n" \
                 "   - Dung lượng gọn nhẹ, chép vào USB mang đi đâu mở lên cũng chạy ngay lập tức mà không cần quyền Administrator.\n\n" \
                 "• Tương thích hoàn hảo: Hỗ trợ Windows 10 và Windows 11 (64-bit)."
    p12_1.font.name = "Arial"
    p12_1.font.size = Pt(11)
    p12_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s12, 6.8, 1.4, 5.7, 5.5, "⚙️ Máy Chủ Nhúng Nội Bộ (Embedded Server)")
    tb12_r = s12.shapes.add_textbox(Inches(6.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf12_r = tb12_r.text_frame
    tf12_r.word_wrap = True
    p12_2 = tf12_r.paragraphs[0]
    p12_2.text = "• Không cần mở Antigravity hay Terminal:\n" \
                 "   - Ứng dụng tích hợp máy chủ Node.js ngầm (embedded-server.cjs) tự động khởi chạy cùng ứng dụng.\n\n" \
                 "• Cơ chế Local Replica SQLite Ngoại Tuyến:\n" \
                 "   - Tự động nạp bản sao cơ sở dữ liệu SQLite cục bộ chứa đầy đủ 29.460 câu hỏi và 10 môn học.\n" \
                 "   - Khi người dùng mang laptop lên máy bay hoặc mất mạng, ứng dụng vẫn hoạt động học tập mượt mà.\n\n" \
                 "• Đăng xuất / Tắt ứng dụng an toàn:\n" \
                 "   - Tự động dọn dẹp port và tiến trình khi người dùng tắt cửa sổ."
    p12_2.font.name = "Arial"
    p12_2.font.size = Pt(11)
    p12_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 13: CƠ SỞ DỮ LIỆU & BẢO MẬT ====================
    s13 = prs.slides.add_slide(blank_layout)
    add_full_background(s13)
    add_slide_header(s13, "Cơ Sở Dữ Liệu Cloudflare D1, R2 & Bảo Mật Tuyệt Đối", slide_num=13)

    add_card(s13, 0.8, 1.4, 5.7, 5.5, "🗄️ 100% Dữ Liệu Thật Trên Cloudflare D1 & R2")
    tb13_l = s13.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf13_l = tb13_l.text_frame
    tf13_l.word_wrap = True
    p13_1 = tf13_l.paragraphs[0]
    p13_1.text = "• Cam kết Zero Mock Data (Cấm bịa đặt thông tin):\n" \
                 "   - 10 Môn học chính khóa đại học CMC\n" \
                 "   - 4.094 Bài học chi tiết\n" \
                 "   - 29.460 Câu hỏi trắc nghiệm có đáp án và giải thích\n\n" \
                 "• Cloudflare R2 Bucket Object Storage:\n" \
                 "   - Lưu trữ Avatar, Cover ảnh bìa người dùng\n" \
                 "   - Lưu trữ tài nguyên hình ảnh Mascot Thỏ Bảy Màu\n\n" \
                 "• Bảng dữ liệu mới chuẩn 3NF:\n" \
                 "   - friendships, friend_messages, user_daily_quests, user_leaderboard_rewards."
    p13_1.font.name = "Arial"
    p13_1.font.size = Pt(11)
    p13_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s13, 6.8, 1.4, 5.7, 5.5, "🔒 Bảo Mật Mã Nguồn & GitHub Protection")
    tb13_r = s13.shapes.add_textbox(Inches(6.95), Inches(2.0), Inches(5.4), Inches(4.7))
    tf13_r = tb13_r.text_frame
    tf13_r.word_wrap = True
    p13_2 = tf13_r.paragraphs[0]
    p13_2.text = "• Không để lộ Secret trên Git:\n" \
                 "   - Toàn bộ Cloudflare API Tokens, Account ID, R2 Access Keys và DeepSeek AI Key được cô lập hoàn toàn trong biến môi trường bí mật .dev.vars (đã được .gitignore bảo vệ).\n\n" \
                 "• Vượt qua GitHub Secret Push Protection:\n" \
                 "   - Mọi commit đều được quét tự động bởi hệ thống bảo vệ của GitHub, đảm bảo tính an toàn bảo mật tuyệt đối cho dự án phần mềm theo tiêu chuẩn ngành công nghiệp."
    p13_2.font.name = "Arial"
    p13_2.font.size = Pt(11)
    p13_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 14: KẾT QUẢ KIỂM THỬ ====================
    s14 = prs.slides.add_slide(blank_layout)
    add_full_background(s14)
    add_slide_header(s14, "Kết Quả Kiểm Thử (QA) & Đánh Giá Đóng Góp Thành Viên", slide_num=14)

    add_card(s14, 0.8, 1.4, 6.0, 5.5, "✅ Kết Quả Kiểm Thử Tính Năng (Pass 100%)")
    tb14_l = s14.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(5.7), Inches(4.7))
    tf14_l = tb14_l.text_frame
    tf14_l.word_wrap = True
    p14_1 = tf14_l.paragraphs[0]
    p14_1.text = "• TC-01: Soi hồ sơ người dùng khác (/api/users/:id) ➔ PASSED\n" \
                 "• TC-02: Nhắn tin trực tiếp & hiển thị Unread Badge ➔ PASSED\n" \
                 "• TC-03: Ôn tập câu sai theo môn, trắc nghiệm tức thì ➔ PASSED\n" \
                 "• TC-04: Đóng băng chuỗi ngày học Duolingo (🧊) ➔ PASSED\n" \
                 "• TC-05: Quá 2 ngày tự động chuyển sang lịch số thường ➔ PASSED\n" \
                 "• TC-06: Mở khóa gói Mascot Học Bá Siêu Đẳng ➔ PASSED\n" \
                 "• TC-07: Đóng gói và chạy file exe trên Windows ➔ PASSED"
    p14_1.font.name = "Arial"
    p14_1.font.size = Pt(11)
    p14_1.font.color.rgb = COLOR_DARK_TEXT

    add_card(s14, 7.1, 1.4, 5.4, 5.5, "🌟 Đánh Giá Nỗ Lực Thành Viên")
    tb14_r = s14.shapes.add_textbox(Inches(7.25), Inches(2.0), Inches(5.1), Inches(4.7))
    tf14_r = tb14_r.text_frame
    tf14_r.word_wrap = True
    p14_2 = tf14_r.paragraphs[0]
    p14_2.text = "🏆 Lê Văn An (Trưởng nhóm):\n" \
                 "   - Kiến trúc hệ thống, API D1/R2, Soi profile, Nhắn tin, Build Desktop\n" \
                 "   - Xếp loại: A+ (Xuất sắc)\n\n" \
                 "🏆 Nguyễn Như Quỳnh:\n" \
                 "   - Thiết kế UI/UX Tone Pink, Giao diện 20 Mascot, Ôn tập câu sai trắc nghiệm\n" \
                 "   - Xếp loại: A+ (Xuất sắc)\n\n" \
                 "🏆 Nguyễn Văn Tú:\n" \
                 "   - Thuật toán Chuỗi Streak Freeze, Quests ngày, Giải đấu tuần, QA & Build\n" \
                 "   - Xếp loại: A+ (Xuất sắc)"
    p14_2.font.name = "Arial"
    p14_2.font.size = Pt(10.5)
    p14_2.font.color.rgb = COLOR_DARK_TEXT

    # ==================== SLIDE 15: LỜI CẢM ƠN ====================
    s15 = prs.slides.add_slide(blank_layout)
    add_full_background(s15)
    add_slide_header(s15, "Lời Tri Ân Giảng Viên & Hỏi Đáp (Q&A)", slide_num=15)

    add_card(s15, 1.5, 1.4, 10.333, 5.5, border_color=COLOR_PRIMARY_PINK, bg_color=COLOR_WHITE)

    tb15 = s15.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(4.7))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p15_1 = tf15.paragraphs[0]
    p15_1.alignment = PP_ALIGN.CENTER
    p15_1.text = "💖 LỜI CẢM ƠN CHÂN THÀNH TỚI CÔ VÂN ANH"
    p15_1.font.name = "Arial"
    p15_1.font.size = Pt(22)
    p15_1.font.bold = True
    p15_1.font.color.rgb = COLOR_DEEP_BERRY

    p15_2 = tf15.add_paragraph()
    p15_2.alignment = PP_ALIGN.CENTER
    p15_2.text = "\nNhóm TONE PINK xin chân thành gửi lời cảm ơn sâu sắc nhất tới Cô Vân Anh — " \
                 "Giảng viên bộ môn Ứng dụng phần mềm.\n" \
                 "Những góp ý và hướng dẫn quý báu của Cô đã giúp dự án 3TG All-in-One hoàn thiện vượt bậc, " \
                 "đạt tiêu chuẩn ứng dụng thương mại thực tế với trải nghiệm người dùng cuốn hút, chuẩn chỉ và đẳng cấp!\n"
    p15_2.font.name = "Arial"
    p15_2.font.size = Pt(12.5)
    p15_2.font.color.rgb = COLOR_DARK_TEXT

    p15_3 = tf15.add_paragraph()
    p15_3.alignment = PP_ALIGN.CENTER
    p15_3.text = "━━━━━━━━━━━━━━━━ 🌸 ━━━━━━━━━━━━━━━━\n" \
                 "MÃ NGUỒN GITHUB: https://github.com/Anniehatani/appstudyAI.git (Nhánh main)\n" \
                 "ỨNG DỤNG ĐÃ SẴN SÀNG CHO PHẦN THUYẾT TRÌNH VÀ DEMO TRỰC TIẾP!"
    p15_3.font.name = "Arial"
    p15_3.font.size = Pt(11.5)
    p15_3.font.bold = True
    p15_3.font.color.rgb = COLOR_NAVY

    output_path = os.path.abspath("tổng kết sprint 2/SLIDE_THUYET_TRINH_SPRINT_2_THO_7_MAU.pptx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print("[SUCCESS] Da tao thanh cong file PowerPoint: SLIDE_THUYET_TRINH_SPRINT_2_THO_7_MAU.pptx")

if __name__ == "__main__":
    create_sprint2_slides()
