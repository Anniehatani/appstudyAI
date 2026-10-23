import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_sprint2_report():
    doc = docx.Document()
    
    # Page setup - Margins: 0.8 inch for spacious modern layout
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = True
        
        # Header for normal pages
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BÁO CÁO TỔNG KẾT SPRINT 2 | ỨNG DỤNG PHẦN MỀM - GV: VÂN ANH")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        # Footer for normal pages
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("3TG Học Vui Mỗi Ngày • Nhóm TONE PINK • Trường Đại học CMC")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(4.5)

    # Palette Constants (Tone Pink & Professional Navy Accents)
    PINK_HEX = "E11D48"       # Primary Accent / Thỏ Bảy Màu Pink #E11D48
    PINK_LIGHT_HEX = "FFF1F2" # Soft Blush Pink fill #FFF1F2
    PINK_BORDER_HEX = "FECDD3"# Card border #FECDD3
    NAVY_HEX = "1E3A8A"       # Primary Headers #1E3A8A
    DARK_TEXT_HEX = "0F172A"  # High contrast text #0F172A
    BG_LIGHT_HEX = "F8FAFC"   # Table row light #F8FAFC
    BORDER_HEX = "CBD5E1"     # Table border #CBD5E1
    GREEN_HEX = "059669"      # Success badge
    PURPLE_HEX = "7C3AED"     # Purple accent

    # Helper functions
    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        for child in list(tcPr):
            if child.tag.endswith('shd'):
                tcPr.remove(child)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=130, bottom=130, left=170, right=170):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_table_borders(table, color=BORDER_HEX, sz="4"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(14.5)
        run.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C) # Rose Pink 700
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(13)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        return p

    def add_callout(title, text, hex_color=PINK_LIGHT_HEX, border_color=PINK_HEX):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, hex_color)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
        
        # Left thick border
        tcPr = cell._tc.get_or_add_tcPr()
        bdr = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'<w:top w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(bdr)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(3)
        trun = cp.add_run(f"📌 {title}\n")
        trun.bold = True
        trun.font.name = 'Calibri'
        trun.font.size = Pt(10.5)
        trun.font.color.rgb = RGBColor(0x9F, 0x12, 0x39)
        
        mrun = cp.add_run(text)
        mrun.font.name = 'Calibri'
        mrun.font.size = Pt(10)
        mrun.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def add_figure(img_path, caption, width_in=5.2):
        abs_img = os.path.abspath(img_path)
        if os.path.exists(abs_img):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(abs_img, width=Inches(width_in))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_c = p_cap.add_run(caption)
            r_c.font.name = 'Calibri'
            r_c.font.size = Pt(9)
            r_c.font.italic = True
            r_c.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # ==================== TRANG BÌA ====================
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pre.paragraph_format.space_before = Pt(24)
    p_pre.paragraph_format.space_after = Pt(4)
    r_uni = p_pre.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO • TRƯỜNG ĐẠI HỌC CMC\nKHOA CÔNG NGHỆ THÔNG TIN")
    r_uni.font.name = "Calibri"
    r_uni.font.size = Pt(11)
    r_uni.bold = True
    r_uni.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Decorative Pink Ribbon Line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(28)
    r_line = p_line.add_run("━━━━━━━━━━━━━━━━ 🌸 ━━━━━━━━━━━━━━━━")
    r_line.font.color.rgb = RGBColor(0xFB, 0x71, 0x85)
    r_line.font.size = Pt(12)

    # Title
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_after = Pt(8)
    r_tit = p_tit.add_run("BÁO CÁO TỔNG KẾT SPRINT 2")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(26)
    r_tit.bold = True
    r_tit.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(28)
    r_sub = p_sub.add_run("HỆ THỐNG MẠNG XÃ HỘI HỌC TẬP, GAME HÓA CHUỖI STREAK DUOLINGO,\nBỘ SƯU TẬP MASCOT THỎ BẢY MÀU & ĐÓNG GÓI ỨNG DỤNG DESKTOP ĐỘC LẬP")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Project Pill Box Table
    tbl_info = doc.add_table(rows=6, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_info, color="FECDD3", sz="6")
    
    info_data = [
        ("Học phần / Môn học:", "Ứng dụng phần mềm (Software Application Development)"),
        ("Giảng viên hướng dẫn:", "Cô Vân Anh"),
        ("Nhóm thực hiện:", "Nhóm TONE PINK"),
        ("Dự án phần mềm:", "3TG — Web & Desktop App Học tập Thông minh phong cách Duolingo"),
        ("Mascot biểu tượng:", "Giáo sư Thỏ Bảy Màu (Độc quyền 20 sắc thái biểu cảm)"),
        ("Thời gian nghiệm thu:", "Tháng 10/2026")
    ]
    for idx, (lbl, val) in enumerate(info_data):
        c0 = tbl_info.cell(idx, 0)
        c1 = tbl_info.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, "FFF1F2" if idx % 2 == 0 else "FFFFFF")
        set_cell_background(c1, "FFF1F2" if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(lbl)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x9F, 0x12, 0x39)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        if idx in [1, 2]:
            r1.bold = True

    # Members Table
    p_mem_head = doc.add_paragraph()
    p_mem_head.paragraph_format.space_before = Pt(24)
    p_mem_head.paragraph_format.space_after = Pt(6)
    r_mh = p_mem_head.add_run("DANH SÁCH THÀNH VIÊN NHÓM TONE PINK:")
    r_mh.bold = True
    r_mh.font.size = Pt(11)
    r_mh.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    tbl_mem = doc.add_table(rows=4, cols=4)
    tbl_mem.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_mem, color="FECDD3", sz="6")

    headers_mem = ["STT", "Họ và Tên", "Vai trò trong Sprint 2", "Phụ trách chuyên môn chính"]
    for c_idx, h_text in enumerate(headers_mem):
        cell = tbl_mem.cell(0, c_idx)
        set_cell_background(cell, "E11D48")
        set_cell_margins(cell, 120, 120, 120, 120)
        cp = cell.paragraphs[0]
        crun = cp.add_run(h_text)
        crun.bold = True
        crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        crun.font.size = Pt(9.5)

    members_data = [
        ("1", "Lê Văn An", "Trưởng nhóm", "Kiến trúc hệ thống, API D1/R2, Mạng xã hội bạn bè & Tin nhắn, Đóng gói Desktop"),
        ("2", "Nguyễn Như Quỳnh", "Thành viên", "Thiết kế UI/UX Tone Pink, Giao diện Mascot 20 biểu cảm, Ôn tập câu sai trắc nghiệm"),
        ("3", "Nguyễn Văn Tú", "Thành viên", "Logic Game hóa Duolingo (Streak Freeze, Quests, Leagues), Kiểm thử tự động & Build Exe")
    ]
    for row_idx, r_data in enumerate(members_data, start=1):
        for col_idx, val in enumerate(r_data):
            cell = tbl_mem.cell(row_idx, col_idx)
            set_cell_background(cell, "FFF7FA" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 90, 90, 110, 110)
            cp = cell.paragraphs[0]
            crun = cp.add_run(val)
            crun.font.size = Pt(9.5)
            if col_idx == 1:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C)

    doc.add_page_break()

    # ==================== TÓM TẮT ĐIỀU HÀNH ====================
    add_h1("TÓM TẮT ĐIỀU HÀNH SPRINT 2 (EXECUTIVE SUMMARY)")
    
    p_sum = doc.add_paragraph()
    p_sum.add_run(
        "Sprint 2 của dự án 3TG All-in-One là giai đoạn mở rộng và hoàn thiện đột phá, chuyển đổi ứng dụng từ "
        "một nền tảng học tập đơn lẻ thành một Mạng xã hội học tập thông minh (Social Learning Platform) hoàn chỉnh, "
        "kết hợp cơ chế Game hóa (Gamification) đỉnh cao chuẩn Duolingo và hình tượng linh vật Giáo sư Thỏ Bảy Màu độc quyền. "
        "Toàn bộ mục tiêu đặt ra cho Sprint 2 đã được Nhóm TONE PINK hoàn thành xuất sắc 100% đúng tiến độ dưới sự hướng dẫn của Giảng viên Vân Anh."
    )

    add_callout(
        "KẾT QUẢ VƯỢT TRỘI ĐẠT ĐƯỢC TRONG SPRINT 2",
        "• 100% Tính năng Mạng xã hội: Soi hồ sơ học viên khác, Kết bạn 1-chạm, Nhắn tin trực tiếp thời gian thực có unread badge.\n"
        "• 100% Cơ chế Chuỗi Streak Duolingo: Giữ chuỗi đóng băng (Streak Freeze) tối đa 2 lần/tuần, tự động chuyển đổi lịch số thường.\n"
        "• Hệ sinh thái 20 Mascot Thỏ Bảy Màu: Chia thành 3 gói (6 con mặc định, Gói Học Bá Siêu Đẳng 7 con, Gói Thánh Hề Bá Đạo 7 con).\n"
        "• Ôn tập câu sai chuẩn đại học: Phân loại theo từng môn, làm lại trắc nghiệm 4 đáp án trực tiếp kèm lời giải chi tiết.\n"
        "• Đóng gói Desktop chuyên nghiệp: Xuất file 3TG Học Vui Mỗi Ngày Setup 1.0.0.exe và Portable chạy độc lập không phụ thuộc môi trường."
    )

    # ==================== PHẦN 1: TỔNG QUAN SPRINT 2 ====================
    add_h1("PHẦN 1: TỔNG QUAN MỤC TIÊU SPRINT 2 & PHÂN CÔNG NHIỆM VỤ")
    
    add_h2("1.1. Mục tiêu cốt lõi của Sprint 2")
    p_obj = doc.add_paragraph()
    p_obj.add_run(
        "Dựa trên phản hồi từ Sprint 1 và yêu cầu thực tiễn của môn học Ứng dụng phần mềm, Nhóm TONE PINK xác định 4 trụ cột phát triển chính trong Sprint 2:\n"
        "1. Trụ cột Tương tác Xã hội (Social Connectivity): Cho phép người học kết nối, xem thành tích của nhau và trao đổi bài vở tức thì.\n"
        "2. Trụ cột Game hóa Sâu sắc (Deep Gamification): Chuỗi ngày học Duolingo với bảo vệ đóng băng, Nhiệm vụ hàng ngày, Giải đấu tuần thưởng Kim Cương.\n"
        "3. Trụ cột Trải nghiệm Người dùng (UX & Mascot Personality): Thay thế hoàn toàn giao diện mặc định bằng Dropdown 3D Tone Pink, tích hợp bộ sưu tập 20 Mascot Thỏ Bảy Màu.\n"
        "4. Trụ cột Độc lập & Đóng gói (Portability & Desktop App): Đóng gói ứng dụng desktop cài đặt hoàn chỉnh cho Windows, chạy mượt mà không cần mở lệnh."
    )

    add_h2("1.2. Bảng phân công chi tiết công việc trong Sprint 2")
    
    tbl_tasks = doc.add_table(rows=7, cols=5)
    tbl_tasks.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_tasks, color="FECDD3", sz="6")

    th_tasks = ["STT", "Hạng mục công việc", "Thành viên phụ trách", "Trọng số", "Trạng thái nghiệm thu"]
    for c_i, th in enumerate(th_tasks):
        cell = tbl_tasks.cell(0, c_i)
        set_cell_background(cell, "E11D48")
        set_cell_margins(cell, 100, 100, 100, 100)
        cp = cell.paragraphs[0]
        crun = cp.add_run(th)
        crun.bold = True
        crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        crun.font.size = Pt(9)

    tasks_data = [
        ("1", "Tính năng Soi hồ sơ (/api/users/:id) & Kết bạn friendships", "Lê Văn An", "15%", "Hoàn thành 100%"),
        ("2", "Nhắn tin trực tiếp thời gian thực, Unread Badge & Empty State", "Lê Văn An", "15%", "Hoàn thành 100%"),
        ("3", "Thiết kế lại Ôn tập câu sai: Trắc nghiệm 4 đáp án theo từng môn", "Nguyễn Như Quỳnh", "15%", "Hoàn thành 100%"),
        ("4", "Cơ chế Hóa đá chuỗi Streak Freeze 2 lần/tuần & Lịch Duolingo", "Nguyễn Văn Tú", "15%", "Hoàn thành 100%"),
        ("5", "Bộ sưu tập 20 Mascot Thỏ Bảy Màu & Gói mở khóa Gacha", "Nguyễn Như Quỳnh", "15%", "Hoàn thành 100%"),
        ("6", "Nhiệm vụ hàng ngày, Giải đấu tuần thưởng Kim Cương & Custom Dropdown 3D", "Lê Văn An & Tú", "25%", "Hoàn thành 100%")
    ]
    for r_i, r_data in enumerate(tasks_data, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl_tasks.cell(r_i, c_i)
            set_cell_background(cell, "FFF7FA" if r_i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)
            cp = cell.paragraphs[0]
            crun = cp.add_run(val)
            crun.font.size = Pt(9)
            if c_i == 4:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

    # ==================== PHẦN 2: CHI TIẾT CÁC TÍNH NĂNG SPRINT 2 ====================
    add_h1("PHẦN 2: CHI TIẾT CÁC TÍNH NĂNG VÀ CẢI TIẾN ĐỘT PHÁ SPRINT 2")

    add_h2("2.1. Tính năng Soi Hồ Sơ (Profile Inspection) & Mạng xã hội Kết Bạn")
    p_p1 = doc.add_paragraph()
    p_p1.add_run(
        "Trong Sprint 2, nhóm TONE PINK đã xây dựng hoàn chỉnh phân hệ Mạng xã hội học tập. Người học giờ đây không còn cảm thấy cô đơn trên hành trình tích lũy kiến thức.\n"
        "• Cơ chế Soi tài khoản: Học viên có thể click vào bất kỳ ai trên Bảng xếp hạng hoặc danh sách bạn bè để xem toàn bộ Hồ sơ học tập cá nhân bao gồm: "
        "Avatar và Ảnh bìa đồng bộ từ Cloudflare R2, Tiểu sử, Trường học, Chuyên ngành, Số bài học đã đạt, Tổng điểm XP tích lũy, Chuỗi Streak hiện tại và Danh hiệu giải đấu.\n"
        "• Cơ chế Kết bạn (Friendship Lifecycle): Tích hợp tính năng gửi lời mời kết bạn, chấp nhận lời mời, hủy kết bạn. Cơ sở dữ liệu Cloudflare D1 bổ sung bảng `friendships` với ràng buộc UNIQUE(user_id, friend_id) chống trùng lặp dữ liệu."
    )

    add_h2("2.2. Nhắn Tin Trực Tiếp Thời Gian Thực (Direct Messaging)")
    p_p2 = doc.add_paragraph()
    p_p2.add_run(
        "Nhóm đã loại bỏ các nút điều hướng thừa thãi để biến Tab Tin Nhắn thành trọng tâm tương tác giữa các học sinh:\n"
        "• Giao diện Chat trực quan: Thanh tab phân đoạn nhỏ gọn, đồng nhất chuẩn Tone Pink. Khi có bạn bè nhắn tin, hệ thống tự động hiển thị bong bóng chat với thời gian gửi rõ ràng (Mine: Hồng Gradient nổi bật, Other: Xám dịu thanh lịch).\n"
        "• Chấm đỏ thông báo (Unread Notification Badge): Khi có tin nhắn mới gửi đến từ bạn bè, Sidebar bên trái và Tab Bạn bè lập tức hiển thị chấm đỏ nhấp nháy (pulsing badge) giúp học viên không bao giờ bỏ lỡ tin nhắn quan trọng.\n"
        "• Empty State Mascot Thỏ Bảy Màu: Nếu chưa có tin nhắn, màn hình hiển thị lời chào dễ thương cùng hình ảnh Thỏ Bảy Màu khích lệ: 'Gửi lời chào để bắt đầu kết nối cùng học tập nào!'."
    )
    add_figure("tổng kết sprint 2/assets/direct_chat.png", "Hình 2.1: Giao diện Nhắn tin trực tiếp giữa các học sinh & Empty State Thỏ Bảy Màu", width_in=5.2)

    add_h2("2.3. Cải Tiến Đột Phá: Phân Hệ Ôn Tập Câu Sai (Mistakes Mastery Review)")
    p_p3 = doc.add_paragraph()
    p_p3.add_run(
        "Ở Sprint 1, chức năng câu sai chỉ đơn thuần đánh dấu đã nhớ hay chưa. Trong Sprint 2, nhóm TONE PINK đã tái cấu trúc 100% luồng ôn tập câu sai chuẩn phương pháp Spaced Repetition (Lặp lại ngắt quãng):\n"
        "• Phân chia câu sai rành mạch theo từng môn học: Người học được lựa chọn môn cần ôn tập (Kỹ năng mềm, Lập trình C++, Web, v.v.), loại bỏ hoàn toàn tình trạng câu sai môn này hiển thị lộn xộn sang môn khác.\n"
        "• Trắc nghiệm 4 đáp án tương tác trực tiếp: Người học phải bấm chọn đáp án A, B, C, D và ấn 'Kiểm tra'. Hệ thống phản hồi ngay lập tức xem lựa chọn đúng hay sai.\n"
        "• Sửa triệt để lỗi kiểm tra đáp án: Khắc phục trạng thái treo 'Đang kiểm tra...', trả về kết quả trong mili-giây cùng phần Lời giải chi tiết (Explanation) giúp sinh viên hiểu rõ bản chất kiến thức.\n"
        "• Cơ chế phục hồi Tim (❤️) và cộng XP: Mỗi câu sai khắc phục thành công sẽ hồi phục 1 Tim và thưởng ngay 5 XP khích lệ."
    )
    add_figure("tổng kết sprint 2/assets/mistakes_review.png", "Hình 2.2: Giao diện làm lại trắc nghiệm câu sai phân loại rành mạch theo từng môn học", width_in=5.2)

    add_h2("2.4. Hệ Thống Chuỗi Streak Duolingo Chuẩn & Cơ Chế Hóa Đá Chuỗi (Streak Freeze)")
    p_p4 = doc.add_paragraph()
    p_p4.add_run(
        "Chuỗi ngày học liên tục (Streak) là yếu tố tâm lý giữ chân người học quan trọng nhất. Sprint 2 mang đến cơ chế đóng băng chuỗi chuẩn xác 100% theo tiêu chuẩn Duolingo toàn cầu:\n"
        "• Giữ chuỗi đóng băng (2 lần/tuần): Mỗi học viên được trang bị tối đa 2 lượt 'Đóng băng chuỗi' mỗi tuần. Nếu bận không vào học 1-2 ngày, chuỗi Streak không bị mất mà ngày đó sẽ chuyển thành biểu tượng tảng băng hóa đá (🧊).\n"
        "• Tự động rã đông khi học tiếp: Ngay khi học viên hoàn thành bài học tiếp theo, chuỗi ngày học được mở khóa và tiếp tục tính lũy tăng dần.\n"
        "• Chuyển đổi lịch số bình thường khi quá hạn: Nếu quá 2 ngày mà người học không tham gia học tập (hết 2 lượt bảo vệ miễn phí), chuỗi streak sẽ bị đứt và lịch học tự động chuyển về dạng lịch số thông thường (normal-date). Để lấy lại chuỗi, người dùng phải vào Cửa hàng dùng Kim Cương mua vật phẩm 'Bảo vệ chuỗi'."
    )
    add_figure("tổng kết sprint 2/assets/streak_freeze.png", "Hình 2.3: Hệ thống Chuỗi ngày học Duolingo với hiệu ứng tảng băng hóa đá 🧊", width_in=5.2)

    add_h2("2.5. Hệ Sinh Thái 20 Mascot Thỏ Bảy Màu & Gói Mở Khóa Gacha (Mascot Packages)")
    p_p5 = doc.add_paragraph()
    p_p5.add_run(
        "Nhóm TONE PINK đã sáng tạo bộ sưu tập độc quyền gồm 20 trạng thái biểu cảm Mascot Thỏ Bảy Màu và chia thành các gói mở khóa phân cấp vô cùng hấp dẫn:\n"
        "• 6 Mascot Cơ bản (Mặc định miễn phí): Thỏ Chăm Chỉ, Thỏ Suy Tư, Thỏ Vui Vẻ, Thỏ Ngủ Gật, Thỏ Quyết Tâm, Thỏ Đọc Sách.\n"
        "• Gói 1 — 'Học Bá Siêu Đẳng' (7 Mascot xịn sò): Mở khóa bằng Kim Cương tích lũy từ học tập, sở hữu hiệu ứng hào quang lấp lánh (Thỏ Tiến Sĩ, Thỏ Bác Học, Thỏ Lập Trình Viên, Thỏ Thủ Khoa, Thỏ Thần Đồng, Thỏ Đạt Giải Nhất, Thỏ Giáo Sư).\n"
        "• Gói 2 — 'Thánh Hề Bá Đạo' (7 Mascot độc lạ): Sở hữu các biểu cảm meme hài hước nổi tiếng của Thỏ Bảy Màu (Thỏ Troll Face, Thỏ Cà Khịa, Thỏ Điệp Viên, Thỏ Siêu Nhân, Thỏ Ninja, Thỏ Bốc Hỏa, Thỏ Đại Gia).\n"
        "• Tối ưu vị trí Mascot: Bỏ linh vật cố định ở góc dưới bên phải (tránh che khuất nội dung bài tập), đưa vào phòng thay đồ linh vật riêng biệt (Mascot Wardrobe)."
    )

    add_h2("2.6. Nhiệm Vụ Hàng Ngày (Daily Quests) & Giải Đấu Tuần Thưởng Kim Cương")
    p_p6 = doc.add_paragraph()
    p_p6.add_run(
        "• Nhiệm vụ hàng ngày (Daily Quests): Hệ thống tự động làm mới vào 00:00 mỗi ngày gồm các thử thách: Khởi động ngày mới (học 1 bài - thưởng 10 💎), Chăm chỉ rèn luyện (học 3 bài - thưởng 15 💎), Chiến thần học tập (học 5 bài - thưởng 20 💎), Điểm số hoàn hảo (bài thi trên 80% - thưởng 15 💎), Khắc phục điểm yếu (ôn tập câu sai - thưởng 10 💎). Học viên nhận Kim Cương ngay lập tức khi hoàn thành.\n"
        "• Giải đấu tuần (Weekly Leaderboards): Thi đua xếp hạng giữa các học viên theo 5 hạng đấu (Đồng, Bạc, Vàng, Lục Bảo, Kim Cương). Vào cuối tuần, TOP 1 nhận ngay 100 💎, TOP 2 nhận 60 💎, TOP 3 nhận 40 💎 kèm Cúp vinh quang và Modal chúc mừng rực rỡ."
    )
    add_figure("tổng kết sprint 2/assets/daily_quests.png", "Hình 2.4: Phân hệ Nhiệm vụ hàng ngày phong phú & Cơ chế nhận thưởng Kim Cương", width_in=5.2)

    add_h2("2.7. Giao Diện Custom Dropdown 3D Tone Pink & Trợ Giảng AI Thông Minh")
    p_p7 = doc.add_paragraph()
    p_p7.add_run(
        "• Loại bỏ thẻ select Windows xấu xí: Nhóm TONE PINK đã tự tay thiết kế và lập trình linh kiện Custom Dropdown Card 3D màu hồng sang trọng (`lumi-dropdown-panel`). Dropdown hỗ trợ animation mở trượt mượt mà, hiệu ứng kính mờ glassmorphism và icon minh họa cho từng phiên hội thoại.\n"
        "• Trợ giảng AI Lumi & Giáo sư Thỏ: Kết nối trực tiếp mô hình ngôn ngữ lớn DeepSeek AI kết hợp kỹ thuật trích xuất ngữ cảnh RAG từ cơ sở dữ liệu học liệu chuẩn trường Đại học CMC, giải đáp mọi bài tập lập trình, toán rời rạc, công nghệ phần mềm."
    )

    add_h2("2.8. Tối Ưu Màn Hình Desktop Tỷ Lệ 150% (Windows Scaling) & Web Mobile")
    p_p8 = doc.add_paragraph()
    p_p8.add_run(
        "• Khắc phục lỗi hiển thị tỷ lệ 150%: Trên các dòng laptop Windows hiện đại có độ phân giải cao (Full HD, 2K) thường mặc định Display Scale 125% - 150%, giao diện cũ trước đây bị tràn lề hoặc khuất nội dung. Nhóm đã tối ưu lại toàn bộ CSS Grid & Flexbox, kích thước container co giãn linh hoạt, đảm bảo hiển thị hoàn hảo ở mọi tỷ lệ.\n"
        "• Cân đối thanh Sidebar & Tab AI: Phân bổ đều khoảng cách, loại bỏ hiện tượng che khuất nút chức năng, giao diện thanh thoát và thẩm mỹ."
    )

    add_h2("2.9. Đóng Gói Ứng Dụng Desktop Độc Lập (Electron Standalone Exe)")
    p_p9 = doc.add_paragraph()
    p_p9.add_run(
        "Người dùng không cần cài đặt Node.js hay gõ lệnh terminal. Nhóm đã đóng gói ứng dụng thành file thực thi Windows chuyên nghiệp:\n"
        "• Bản cài đặt NSIS: `dist-electron/3TG Học Vui Mỗi Ngày Setup 1.0.0.exe` tự động tạo shortcut ngoài Desktop và Start Menu.\n"
        "• Bản Portable: `dist-electron/3TG Học Vui Mỗi Ngày 1.0.0.exe` tải về là chạy ngay không cần cài đặt.\n"
        "• Tích hợp máy chủ nhúng nội bộ (`embedded-server.cjs`) cùng cơ chế Local Replica SQLite, đảm bảo ứng dụng vẫn khởi động và học tập mượt mà ngay cả khi mất mạng internet."
    )

    # ==================== PHẦN 3: KIẾN TRÚC DỮ LIỆU & BẢO MẬT ====================
    add_h1("PHẦN 3: KIẾN TRÚC DỮ LIỆU CLOUDFLARE D1, R2 & BẢO MẬT NGUỒN")

    add_h2("3.1. Các bảng cơ sở dữ liệu mới phát triển trong Sprint 2")
    
    tbl_d1 = doc.add_table(rows=6, cols=3)
    tbl_d1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_d1, color="FECDD3", sz="6")

    th_d1 = ["Tên Bảng D1", "Cấu trúc khóa & Ràng buộc", "Mục đích sử dụng trong Sprint 2"]
    for c_i, th in enumerate(th_d1):
        cell = tbl_d1.cell(0, c_i)
        set_cell_background(cell, "E11D48")
        set_cell_margins(cell, 100, 100, 100, 100)
        cp = cell.paragraphs[0]
        crun = cp.add_run(th)
        crun.bold = True
        crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        crun.font.size = Pt(9.5)

    d1_data = [
        ("friendships", "id PK, user_id, friend_id, status, created_at (UNIQUE)", "Quản lý quan hệ bạn bè, lời mời kết bạn giữa các học viên"),
        ("friend_messages", "id PK, sender_id, receiver_id, message, is_read, created_at", "Lưu trữ tin nhắn trực tiếp thời gian thực, quản lý cờ chưa đọc"),
        ("ai_chats & ai_messages", "chat_id, user_id, subject_id, role, content", "Lưu trữ lịch sử đàm thoại cùng Trợ giảng AI theo từng chủ đề"),
        ("user_daily_quests", "id PK, user_id, quest_date, quest_code, claimed_at", "Ghi nhận trạng thái hoàn thành và nhận thưởng nhiệm vụ ngày"),
        ("user_leaderboard_rewards", "id PK, user_id, week_start_date, rank, gems, claimed_at", "Ghi nhận phần thưởng Kim Cương giải đấu tuần cho TOP 1, 2, 3")
    ]
    for r_i, r_data in enumerate(d1_data, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl_d1.cell(r_i, c_i)
            set_cell_background(cell, "FFF7FA" if r_i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)
            cp = cell.paragraphs[0]
            crun = cp.add_run(val)
            crun.font.size = Pt(9.5)
            if c_i == 0:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C)

    add_h2("3.2. Đảm bảo dữ liệu thật 100% & Bảo mật mã nguồn GitHub")
    p_sec = doc.add_paragraph()
    p_sec.add_run(
        "• Toàn bộ dữ liệu 100% trên Cloudflare D1 & R2: Cam kết tuyệt đối không bịa đặt dữ liệu (Zero Mock Data). Hệ thống quản lý thực tế 29.460 câu hỏi trắc nghiệm, 4.094 bài học, 10 môn học chính khóa Đại học CMC.\n"
        "• Tuân thủ nghiêm ngặt bảo mật Git: Nhóm đã loại bỏ toàn bộ các API Key, Cloudflare Tokens và Secrets ra khỏi các file code nguồn, chuyển vào biến môi trường an toàn `.dev.vars`, vượt qua kiểm duyệt bảo vệ tự động của GitHub Secret Push Protection."
    )

    # ==================== PHẦN 4: KẾT QUẢ KIỂM THỬ SPRINT 2 ====================
    add_h1("PHẦN 4: KẾT QUẢ KIỂM THỬ TÍNH NĂNG SPRINT 2 (TEST RESULTS)")

    tbl_test = doc.add_table(rows=8, cols=4)
    tbl_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_test, color="FECDD3", sz="6")

    th_test = ["Mã Test", "Kịch bản kiểm thử (Test Scenario)", "Kết quả thực tế", "Đánh giá"]
    for c_i, th in enumerate(th_test):
        cell = tbl_test.cell(0, c_i)
        set_cell_background(cell, "E11D48")
        set_cell_margins(cell, 100, 100, 100, 100)
        cp = cell.paragraphs[0]
        crun = cp.add_run(th)
        crun.bold = True
        crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        crun.font.size = Pt(9.5)

    test_data = [
        ("TC-01", "Soi hồ sơ bạn học khác từ Bảng xếp hạng", "Hiển thị đầy đủ avatar R2, XP, Streak, Môn học hoàn thành", "PASSED"),
        ("TC-02", "Gửi tin nhắn trực tiếp giữa 2 tài khoản học viên", "Tin nhắn hiển thị ngay lập tức, có chấm đỏ unread badge", "PASSED"),
        ("TC-03", "Làm lại câu trắc nghiệm sai của môn Lập trình Web", "Phân loại chuẩn môn Web, chọn đáp án đúng phản hồi tức thì kèm giải thích", "PASSED"),
        ("TC-04", "Kiểm tra cơ chế đóng băng chuỗi (Streak Freeze)", "Nghỉ học 1 ngày hiển thị icon tảng băng 🧊, học lại chuỗi không bị mất", "PASSED"),
        ("TC-05", "Nghỉ học quá 2 ngày liên tiếp khi hết lượt bảo vệ", "Lịch học tự động chuyển về dạng lịch số thường chuẩn quy tắc", "PASSED"),
        ("TC-06", "Mở khóa gói Mascot Học Bá Siêu Đẳng trong Cửa hàng", "Trừ Kim Cương chính xác, mở khóa 7 mascot có hào quang lấp lánh", "PASSED"),
        ("TC-07", "Cài đặt và chạy file exe đóng gói trên máy tính khác", "Chạy độc lập mượt mà không cần cài thêm Node.js hay mở IDE", "PASSED")
    ]
    for r_i, r_data in enumerate(test_data, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl_test.cell(r_i, c_i)
            set_cell_background(cell, "FFF7FA" if r_i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)
            cp = cell.paragraphs[0]
            crun = cp.add_run(val)
            crun.font.size = Pt(9.5)
            if c_i == 3:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

    # ==================== PHẦN 5: TỔNG KẾT & ĐÁNH GIÁ ====================
    add_h1("PHẦN 5: ĐÁNH GIÁ ĐÓNG GÓP THÀNH VIÊN & LỜI CẢM ƠN")

    add_h2("5.1. Đánh giá mức độ hoàn thành của các thành viên Nhóm TONE PINK")
    
    tbl_eval = doc.add_table(rows=4, cols=4)
    tbl_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_eval, color="FECDD3", sz="6")

    th_eval = ["Thành viên", "Trách nhiệm chính trong Sprint 2", "Đánh giá nỗ lực", "Xếp loại"]
    for c_i, th in enumerate(th_eval):
        cell = tbl_eval.cell(0, c_i)
        set_cell_background(cell, "E11D48")
        set_cell_margins(cell, 100, 100, 100, 100)
        cp = cell.paragraphs[0]
        crun = cp.add_run(th)
        crun.bold = True
        crun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        crun.font.size = Pt(9.5)

    eval_data = [
        ("Lê Văn An (Trưởng nhóm)", "Dẫn dắt tiến độ, thiết kế API mạng xã hội bạn bè, tin nhắn, đóng gói Electron & bảo mật", "100% nhiệm vụ xuất sắc", "A+ (Xuất sắc)"),
        ("Nguyễn Như Quỳnh", "Thiết kế UI/UX Tone Pink, tái cấu trúc trắc nghiệm câu sai phân môn, bộ sưu tập Mascot 20 biểu cảm", "100% nhiệm vụ xuất sắc", "A+ (Xuất sắc)"),
        ("Nguyễn Văn Tú", "Lập trình logic đóng băng chuỗi Duolingo, nhiệm vụ ngày, giải đấu tuần, chạy test & đóng gói NSIS", "100% nhiệm vụ xuất sắc", "A+ (Xuất sắc)")
    ]
    for r_i, r_data in enumerate(eval_data, start=1):
        for c_i, val in enumerate(r_data):
            cell = tbl_eval.cell(r_i, c_i)
            set_cell_background(cell, "FFF7FA" if r_i % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)
            cp = cell.paragraphs[0]
            crun = cp.add_run(val)
            crun.font.size = Pt(9.5)
            if c_i == 0:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0xBE, 0x12, 0x3C)
            if c_i == 3:
                crun.bold = True
                crun.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

    add_h2("5.2. Lời cảm ơn gửi tới Giảng viên hướng dẫn")
    p_thanks = doc.add_paragraph()
    p_thanks.add_run(
        "Nhóm TONE PINK xin chân thành gửi lời cảm ơn sâu sắc nhất tới Cô Vân Anh — Giảng viên bộ môn Ứng dụng phần mềm. "
        "Những định hướng chuyên môn tận tình và các nhận xét sâu sắc của Cô trong suốt quá trình phát triển dự án là kim chỉ nam "
        "giúp nhóm hoàn thiện sản phẩm đạt chuẩn chất lượng thực tế của một phần mềm thương mại hiện đại.\n\n"
        "Nhóm TONE PINK kính chúc Cô luôn dồi dào sức khỏe, niềm vui và tiếp tục truyền lửa đam mê công nghệ cho các thế hệ sinh viên!"
    )

    output_path = os.path.abspath("tổng kết sprint 2/BÁO CÁO TỔNG KẾT SPRINT 2 - APP 3TG.docx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print("[SUCCESS] Da tao thanh cong file Word: BAO CAO TONG KET SPRINT 2 - APP 3TG.docx")

if __name__ == "__main__":
    create_sprint2_report()
